"""
Lee planilla/contenido_alfa.xlsx, la valida y escribe los JSON en what-the-ship-remembers/datos/.

Uso (desde la raíz del repo):  python herramientas/planilla_a_json.py
Requiere: pip install openpyxl

Si hay ERRORES no se escriben los JSON. Las ADVERTENCIAS no frenan la exportación.
"""
import json
import sys
from pathlib import Path
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
PLANILLA = RAIZ / "planilla" / "contenido_alfa.xlsx"
SALIDA = RAIZ / "what-the-ship-remembers" / "datos"

TIEMPOS = {"pasado", "presente", "futuro", "anomalo"}
CAMPOS_COND = {"Usado", "Activo"}
# Cadena de la pinza y el panel: estos objetos y todo lo que los activa deben tener LuzM <= 2
CADENA_PINZA = ["COM_O", "MAN_B"]
# Salas que necesitan el arreglo del panel (evento arreglo_panel) para llegar a luz máxima
SALAS_PANEL = ("MAN", "COM")
# Objetos clave (⭐): deben poder usarse (o verse, si son forzados) dentro de los 12 turnos
CLAVE = ["MAN_A", "MAN_B", "COM_U", "MAN_N4", "MAN_N5", "COM_E", "COM_H", "COM_J", "COM_L", "COM_R", "COM_O", "COM_S",
         "COM_Q", "SCO_D", "SCO_J", "SCO_O", "PTA_A"]
# Orden de la pinza (21/10/2026): T6 presente "pensé que la dejé acá" → T7 futuro se agarra →
# T9 pasado está otra vez (investigarla es opcional) → T10 presente: diálogo obligatorio del bolsillo
# (solo si no se investigó en el pasado).
ORDEN_PINZA = ["COM_R", "COM_O", "COM_S"]
# Interacciones opcionales que la simulación del recorrido mínimo NO hace (investigar la pinza del pasado
# cancela el diálogo obligatorio del presente).
OPCIONALES_SIM = {"COM_M"}
# Textos que usa el código directamente (no un objeto)
TEXTOS_SISTEMA = ["SIS_SIN_ENERGIA", "SIS_DESCANSAR_PREGUNTA", "SIS_DESCANSAR_SI", "SIS_NORA_MENU",
                  "T_CIERRE_PINZA", "T_CIERRE_GENERAL", "T_CIERRE_FIN"]  # cierre del Alfa (juego.gd, paso 9)
# NORA_OTRA = NORA hablando con datos de otro tiempo (se ve/oye distinto, ver 01 sección 2.1)
HABLANTES = ["PERSONAJE", "NORA", "NORA_OTRA"]
# Hojas de arte para los sprites (tamaño en píxeles) y opciones que se pueden escribir después del nombre
HOJAS_ARTE = {"cosas": (1056, 576), "cosas2": (480, 1104), "pantallas": (480, 576), "puerta": (48, 384)}
OPCIONES_SPRITE = {"girado", "espejo", "repetir", "doble", "pared"}

errores, avisos = [], []
def err(m): errores.append(m)
def aviso(m): avisos.append(m)


def leer_hoja(wb, nombre):
    ws = wb[nombre]
    filas = list(ws.iter_rows(values_only=True))
    cab = [str(c).strip() if c is not None else "" for c in filas[0]]
    out = []
    for i, f in enumerate(filas[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in f):
            continue
        d = {}
        for k, v in zip(cab, f):
            if not k:
                continue
            if isinstance(v, str):
                v = v.strip()
                if v == "":
                    v = None
            d[k] = v
        d["_fila"] = i
        out.append(d)
    return out


def a_int(v, donde, campo):
    try:
        if v is None:
            raise ValueError
        f = float(v)
        if f != int(f):
            raise ValueError
        return int(f)
    except (ValueError, TypeError):
        err(f"{donde}: '{campo}' tiene que ser un número entero (vale {v!r})")
        return None


def si_no(v):
    return str(v).strip().lower() in ("sí", "si", "true", "1", "x") if v is not None else False


def main():
    if not PLANILLA.exists():
        print(f"No encuentro la planilla en {PLANILLA}")
        return 2
    wb = load_workbook(PLANILLA, data_only=True)

    # ---------- tablas auxiliares ----------
    salas = {}
    for r in leer_hoja(wb, "salas"):
        sid = r["sala"]
        if sid in salas:
            err(f"salas fila {r['_fila']}: sala '{sid}' repetida")
        salas[sid] = {
            "numero": a_int(r["numero"], f"salas {sid}", "numero"),
            "nombre": r["nombre"],
            "desbloqueada_inicio": si_no(r["desbloqueada_inicio"]),
            "luz_inicial": a_int(r["luz_inicial"], f"salas {sid}", "luz_inicial"),
            "luz_max_sin_arreglo": a_int(r["luz_max_sin_arreglo"], f"salas {sid}", "luz_max_sin_arreglo"),
            "luz_max_alfa": a_int(r["luz_max_alfa"], f"salas {sid}", "luz_max_alfa"),
            "capacidad_a_luz_max": a_int(r["capacidad_a_luz_max"], f"salas {sid}", "capacidad_a_luz_max"),
        }

    textos = {}
    for r in leer_hoja(wb, "textos"):
        tid = r["id"]
        if tid is None:
            err(f"textos fila {r['_fila']}: falta id"); continue
        if not r.get("texto"):
            err(f"textos fila {r['_fila']}: '{tid}' sin texto")
        if r.get("hablante") not in HABLANTES:
            err(f"textos fila {r['_fila']}: '{tid}' hablante '{r.get('hablante')}' no válido (usar {', '.join(HABLANTES)})")
        textos.setdefault(tid, []).append({"orden": a_int(r["orden"], f"textos {tid}", "orden"),
                                           "hablante": r["hablante"], "texto": r["texto"]})
    for tid, lineas in textos.items():
        ords = [l["orden"] for l in lineas]
        if len(set(ords)) != len(ords):
            err(f"textos: '{tid}' tiene números de orden repetidos")
        lineas.sort(key=lambda l: l["orden"] or 0)

    def tabla_simple(nombre, campos):
        out = {}
        for r in leer_hoja(wb, nombre):
            if r["id"] in out:
                err(f"{nombre} fila {r['_fila']}: id '{r['id']}' repetido")
            out[r["id"]] = {c: r.get(c) for c in campos}
        return out

    anotaciones = {k: v["texto"] for k, v in tabla_simple("anotaciones", ["texto"]).items()}
    items = tabla_simple("items", ["nombre", "descripcion"])
    audios = tabla_simple("audios", ["descripcion"])
    eventos = tabla_simple("eventos", ["descripcion"])
    config = {r["clave"]: r["valor"] for r in leer_hoja(wb, "config")}

    # ---------- sprites (catálogo de recortes de las hojas de arte) ----------
    sprites = {}
    if "sprites" in wb.sheetnames:
        for r in leer_hoja(wb, "sprites"):
            donde = f"sprites fila {r['_fila']}"
            n, hoja = r.get("nombre"), r.get("hoja")
            if not n:
                err(f"{donde}: falta el nombre"); continue
            if n in sprites or n == "ninguno":
                err(f"{donde}: nombre '{n}' repetido o reservado")
            if hoja not in HOJAS_ARTE:
                err(f"{donde}: hoja '{hoja}' no existe (opciones: {', '.join(HOJAS_ARTE)})"); continue
            x, y, w, h = (a_int(r.get(k), donde, k) for k in ("x", "y", "ancho", "alto"))
            if None in (x, y, w, h):
                continue
            W, H = HOJAS_ARTE[hoja]
            if x < 0 or y < 0 or w <= 0 or h <= 0 or x + w > W or y + h > H:
                err(f"{donde}: el recorte {x},{y} {w}x{h} se sale de la hoja {hoja} ({W}x{H})")
            sprites[n] = {"hoja": hoja, "region": [x, y, w, h]}

    def revisar_sprite(valor, donde, campo):
        if valor is None:
            return None
        partes = str(valor).split()
        if partes[0] != "ninguno" and partes[0] not in sprites:
            err(f"{donde}: {campo} '{partes[0]}' no está en la hoja sprites")
        for o in partes[1:]:
            if o not in OPCIONES_SPRITE:
                err(f"{donde}: {campo} tiene una opción desconocida '{o}' (opciones: {', '.join(sorted(OPCIONES_SPRITE))})")
        return " ".join(partes)

    # ---------- guion ----------
    guion = {}
    for r in leer_hoja(wb, "guion"):
        t = a_int(r["turno"], "guion", "turno")
        guion[t] = {}
        for s in salas:
            v = r.get(s)
            if v not in TIEMPOS:
                err(f"guion turno {t}: la sala {s} tiene tiempo inválido ({v!r})")
            guion[t][s] = v
        if guion[t].get("MAN") != "presente":
            err(f"guion turno {t}: la sala de mando tiene que estar en presente")
    n_turnos = int(config.get("turnos_alfa", 12))
    if sorted(guion) != list(range(1, n_turnos + 1)):
        err(f"guion: tiene que tener los turnos 1 a {n_turnos}, uno por fila")

    # ---------- conexiones ----------
    conexiones = []
    for r in leer_hoja(wb, "conexiones"):
        a, b = r["desde"], r["hasta"]
        for x in (a, b):
            if x not in salas:
                err(f"conexiones fila {r['_fila']}: la sala '{x}' no existe")
        if a == b:
            err(f"conexiones fila {r['_fila']}: conecta una sala consigo misma")
        conexiones.append([a, b, si_no(r.get("vuelta_gratis"))])
    for s in salas:
        if not any(s in c[:2] for c in conexiones):
            err(f"conexiones: la sala {s} no está conectada con ninguna")

    # ---------- objetos ----------
    objetos = {}
    for r in leer_hoja(wb, "objetos"):
        donde = f"objetos fila {r['_fila']}"
        sala, letra = r["sala"], r["letra"]
        if sala not in salas:
            err(f"{donde}: sala '{sala}' no existe en la hoja salas"); continue
        if letra is None:
            err(f"{donde}: falta la letra"); continue
        oid = f"{sala}_{letra}"
        if r.get("id") and r["id"] != oid:
            err(f"{donde}: el id '{r['id']}' no coincide con sala+letra ({oid})")
        if oid in objetos:
            err(f"{donde}: id '{oid}' repetido (la letra tiene que ser única en la sala)")
        donde = f"objeto {oid}"
        tiempo = r["tiempo"]
        if tiempo not in TIEMPOS:
            err(f"{donde}: tiempo inválido ({tiempo!r}); tiene que ser un único valor: pasado, presente o futuro")
        luz = a_int(r["luzM"], donde, "luzM")
        if luz is not None and luz not in (1, 2, 3):
            err(f"{donde}: luzM tiene que ser 1, 2 o 3")
        costo = a_int(r["costo"], donde, "costo")
        turno = a_int(r["turnoM"], donde, "turnoM")
        conds = []
        for k in ("cond1", "cond2", "cond3"):
            c = r.get(k)
            if c is None:
                continue
            partes = str(c).split(".")
            if len(partes) != 2 or partes[1] not in CAMPOS_COND:
                err(f"{donde}: condición '{c}' mal escrita (formato: ID.Usado o ID.Activo)"); continue
            conds.append({"objeto": partes[0], "campo": partes[1].lower()})
        tipo = r.get("tipoCond")
        if conds and tipo not in ("Y", "O"):
            err(f"{donde}: tiene condiciones pero tipoCond no es Y ni O")
        disps = []
        for i in (1, 2, 3, 4):
            d, p = r.get(f"disp{i}"), r.get(f"param{i}")
            if d is None and p is None:
                continue
            d = a_int(d, donde, f"disp{i}")
            if d not in (0, 1, 2, 3, 4, 5):
                err(f"{donde}: disp{i} tiene tipo inválido ({d})"); continue
            if p is None:
                err(f"{donde}: disp{i} sin parámetro"); continue
            disps.append({"tipo": d, "param": str(p)})
        objetos[oid] = {
            "sala": sala, "letra": str(letra), "descripcion": r.get("descripcion"),
            "tiempo": tiempo, "turno_min": turno, "luz_min": luz, "costo": costo,
            "forzada": si_no(r.get("forzada")), "texto_al_entrar": r.get("texto_al_entrar"),
            "tipo_cond": tipo if conds else None, "condiciones": conds, "disparadores": disps,
            "pos": [a_int(r.get(k), donde, k) for k in ("x", "y", "ancho", "alto")],
            "nombre": r.get("nombre"),
            "turno_max": a_int(r["turnoMax"], donde, "turnoMax") if r.get("turnoMax") is not None else None,
            "oculto": si_no(r.get("oculto")),
            "sprite": revisar_sprite(r.get("sprite"), donde, "sprite"),
            "sprite_usado": revisar_sprite(r.get("sprite_usado"), donde, "sprite_usado"),
        }
        o = objetos[oid]
        if o["turno_max"] is not None and turno is not None and o["turno_max"] < turno:
            err(f"{donde}: turnoMax ({o['turno_max']}) es menor que turnoM ({turno})")
        if o["oculto"] and not o["forzada"]:
            err(f"{donde}: es oculto pero no es forzada (nunca se podría usar)")

    # ---------- muebles (fijos en todos los tiempos) ----------
    muebles = []
    for r in leer_hoja(wb, "muebles"):
        donde = f"muebles fila {r['_fila']}"
        if r["sala"] not in salas:
            err(f"{donde}: sala '{r['sala']}' no existe"); continue
        if not r.get("nombre"):
            err(f"{donde}: falta el nombre")
        en_futuro = r.get("en_futuro")
        if en_futuro not in (None, "caído", "roto"):
            err(f"{donde}: en_futuro tiene que ser caído, roto o vacío (vale {en_futuro!r})")
        muebles.append({"sala": r["sala"], "nombre": r.get("nombre"),
                        "pos": [a_int(r.get(k), donde, k) for k in ("x", "y", "ancho", "alto")],
                        "solido": si_no(r.get("solido")), "en_futuro": en_futuro,
                        "sprite": revisar_sprite(r.get("sprite"), donde, "sprite"),
                        "sprite_futuro": revisar_sprite(r.get("sprite_futuro"), donde, "sprite_futuro")})

    # ---------- posiciones (tiles dentro de la sala) ----------
    TAM_SALAS = {"MAN": (10, 6), "COM": (10, 6), "SCO": (12, 6), "PTA": (6, 4)}  # escenas/salas/*.tscn
    cosas = [(f"objeto {oid}", o) for oid, o in objetos.items()] + [(f"mueble {m['nombre']}", m) for m in muebles]
    for quien, o in cosas:
        x, y, w, h = o["pos"]
        if None in (x, y, w, h):
            continue
        if w < 1 or h < 1:
            err(f"{quien}: ancho y alto tienen que ser al menos 1")
        tam = TAM_SALAS.get(o["sala"])
        if tam and (x < 0 or y < 0 or x + w > tam[0] or y + h > tam[1]):
            err(f"{quien}: la posición ({x},{y},{w}x{h}) se sale de la sala {o['sala']} ({tam[0]}x{tam[1]} tiles)")

    # ---------- referencias ----------
    usados = {"textos": set(), "anotaciones": set(), "items": set(), "audios": set(), "eventos": set()}
    tablas = {0: ("eventos", eventos), 1: ("textos", textos), 2: ("items", items),
              3: ("audios", audios), 4: ("anotaciones", anotaciones)}
    for oid, o in objetos.items():
        for c in o["condiciones"]:
            if c["objeto"] not in objetos:
                err(f"objeto {oid}: la condición apunta a '{c['objeto']}', que no existe")
            elif c["objeto"] == oid:
                err(f"objeto {oid}: se tiene a sí mismo como condición")
        for d in o["disparadores"]:
            if d["tipo"] == 5:
                if d["param"] not in objetos:
                    err(f"objeto {oid}: desdibujar apunta a '{d['param']}', que no existe")
            else:
                nombre, tabla = tablas[d["tipo"]]
                if d["param"] not in tabla:
                    err(f"objeto {oid}: {nombre} '{d['param']}' no existe en la hoja {nombre}")
                usados[nombre].add(d["param"])
        if o["texto_al_entrar"]:
            if o["texto_al_entrar"] not in textos:
                err(f"objeto {oid}: texto_al_entrar '{o['texto_al_entrar']}' no existe")
            usados["textos"].add(o["texto_al_entrar"])
        if o["disparadores"] and not o["nombre"]:
            err(f"objeto {oid}: falta el nombre (se muestra al acercarse: 'E: investigar <nombre>')")
        if o["forzada"] and o["costo"] not in (0, None):
            err(f"objeto {oid}: es forzada, el costo tiene que ser 0")
        if o["forzada"] and not o["disparadores"]:
            err(f"objeto {oid}: es forzada pero no tiene disparadores")
        if o["turno_min"] is not None and not 1 <= o["turno_min"] <= n_turnos:
            aviso(f"objeto {oid}: turnoM {o['turno_min']} está fuera de los turnos del Alfa")
    for tid in TEXTOS_SISTEMA:
        if tid not in textos:
            err(f"textos: falta '{tid}' (lo usa el código)")
        usados["textos"].add(tid)
    for nombre, tabla in (("textos", textos), ("anotaciones", anotaciones), ("items", items),
                          ("audios", audios), ("eventos", eventos)):
        for k in tabla:
            if k not in usados[nombre]:
                aviso(f"{nombre}: '{k}' no lo usa ningún objeto")

    # ---------- luz de la cadena de la pinza ----------
    def ancestros(oid, vistos):
        for c in objetos.get(oid, {}).get("condiciones", []):
            if c["objeto"] not in vistos:
                vistos.add(c["objeto"])
                ancestros(c["objeto"], vistos)
        return vistos
    cadena = set()
    for oid in CADENA_PINZA:
        if oid in objetos:
            cadena |= {oid} | ancestros(oid, set())
    for oid in sorted(cadena):
        if (objetos[oid]["luz_min"] or 0) > 2:
            err(f"cadena de la pinza: {oid} tiene luzM {objetos[oid]['luz_min']} (tiene que ser <= 2)")

    # ---------- luz posible ----------
    for oid, o in objetos.items():
        s = salas[o["sala"]]
        if o["luz_min"] and s["luz_max_alfa"] is not None and o["luz_min"] > s["luz_max_alfa"]:
            aviso(f"objeto {oid}: luzM {o['luz_min']} pero la sala {o['sala']} llega a {s['luz_max_alfa']} en el Alfa (no se ve; ok si es para la Beta)")

    if errores:
        return reporte(None)

    # ---------- simulación de alcance (sin energía, luz al máximo posible) ----------
    est = {oid: {"activo": False, "usado": False, "desdibujado": False} for oid in objetos}
    abiertas = {s for s, v in salas.items() if v["desbloqueada_inicio"]}
    arreglo = False
    primer_uso, primera_vista = {}, {}

    def cumple(o):
        if not o["condiciones"]:
            return True
        res = [est[c["objeto"]][c["campo"]] if c["campo"] == "activo" else est[c["objeto"]]["usado"]
               for c in o["condiciones"]]
        return all(res) if o["tipo_cond"] == "Y" else any(res)

    def revisar(t):
        cambio = False
        for oid, o in objetos.items():
            if not est[oid]["activo"] and (o["turno_min"] or 1) <= t and cumple(o):
                est[oid]["activo"] = True
                cambio = True
        return cambio

    def luz_sala(s):
        v = salas[s]
        return v["luz_max_alfa"] if (s not in SALAS_PANEL or arreglo) else v["luz_max_sin_arreglo"]

    for t in range(1, n_turnos + 1):
        revisar(t)
        hubo = True
        while hubo:
            hubo = False
            for oid, o in objetos.items():
                e = est[oid]
                if o["sala"] not in abiertas or guion[t][o["sala"]] != o["tiempo"]:
                    continue
                if o["turno_max"] is not None and t > o["turno_max"]:
                    continue
                if not e["activo"] or e["desdibujado"] or luz_sala(o["sala"]) < o["luz_min"]:
                    continue
                primera_vista.setdefault(oid, t)
                if e["usado"] or not o["disparadores"] or oid in OPCIONALES_SIM:
                    continue
                e["usado"] = True
                primer_uso[oid] = t
                for d in o["disparadores"]:
                    if d["tipo"] == 5:
                        est[d["param"]]["desdibujado"] = True
                    elif d["tipo"] == 0 and d["param"] == "arreglo_panel":
                        arreglo = True
                    elif d["tipo"] == 0 and d["param"].startswith("desbloquear_"):
                        abiertas.add(d["param"].split("_", 1)[1])
                revisar(t)
                hubo = True

    for oid in CLAVE:
        if oid not in objetos:
            err(f"objeto clave {oid} no existe en la planilla")
        elif oid not in primer_uso:
            err(f"objeto clave {oid} nunca se puede usar en los {n_turnos} turnos (revisar condiciones, tiempo y guion)")
    for oid, o in objetos.items():
        if oid not in primera_vista and (o["luz_min"] or 0) <= (salas[o["sala"]]["luz_max_alfa"] or 0):
            aviso(f"objeto {oid}: nunca llega a verse en el Alfa (tiempo {o['tiempo']}, condiciones o guion)")
        elif oid in primera_vista and o["disparadores"] and oid not in primer_uso and oid not in OPCIONALES_SIM:
            aviso(f"objeto {oid}: se ve pero nunca se puede usar")
    # La pinza: cada paso en un turno posterior al anterior
    pasos = [(oid, primer_uso.get(oid)) for oid in ORDEN_PINZA]
    for (a, ta), (b, tb) in zip(pasos, pasos[1:]):
        if ta and tb and tb <= ta:
            err(f"pinza: {b} (T{tb}) tiene que pasar después de {a} (T{ta})")

    return reporte({"uso": primer_uso, "vista": primera_vista, "objetos": objetos,
                    "datos": (objetos, salas, guion, conexiones, textos, anotaciones, items, audios, eventos, config, muebles, sprites)})


def reporte(sim):
    print("=" * 60)
    print("VALIDACIÓN DE LA PLANILLA DEL ALFA")
    print("=" * 60)
    if sim:
        cad = ["SCO_D", "SCO_J", "SCO_O", "MAN_N4", "MAN_A", "COM_U", "COM_E", "COM_H", "COM_L", "COM_R", "COM_O", "MAN_B", "COM_M", "COM_S"]
        print("\nRecorrido mínimo (simulación sin energía, en qué turno se puede usar/ver cada cosa):")
        for oid in cad:
            u = sim["uso"].get(oid)
            v = sim["vista"].get(oid)
            print(f"  {oid:7} {sim['objetos'][oid]['descripcion'][:50]:50}  " +
                  (f"uso T{u}" if u else (f"visible T{v}" if v else "NUNCA")))
    print(f"\nErrores: {len(errores)}")
    for e in errores:
        print("  ✗", e)
    print(f"Advertencias: {len(avisos)}")
    for a in avisos:
        print("  !", a)
    if errores:
        print("\nNo se escribieron los JSON. Corregí los errores y volvé a correrlo.")
        return 1
    objetos, salas, guion, conexiones, textos, anotaciones, items, audios, eventos, config, muebles, sprites = sim["datos"]
    SALIDA.mkdir(parents=True, exist_ok=True)
    archivos = {"objetos": objetos, "salas": salas, "guion": {str(k): v for k, v in guion.items()}, "conexiones": conexiones,
                "textos": textos, "anotaciones": anotaciones, "items": items, "audios": audios,
                "eventos": eventos, "config": config, "muebles": muebles, "sprites": sprites}
    for nombre, data in archivos.items():
        (SALIDA / f"{nombre}.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nOK: {len(archivos)} archivos JSON escritos en {SALIDA.relative_to(RAIZ)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
